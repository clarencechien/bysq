"""Build the annual JST (Jordà-Schularick-Taylor Macrohistory, R6) panel
used by the v4 annual engine. Keeps only what the engine needs:
  year, iso, eq_tr, bond_tr, bill_rate, cpi, xrusd, gdp, pop, crisisJST
Real returns are NOT precomputed here (the engine deflates with cpi so that
the return/inflation covariance survives block sampling).
Run: python data/build_jst.py  -> data/processed/jst_annual.csv"""
import csv, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
COLS = ["year", "iso", "eq_tr", "bond_tr", "bill_rate", "cpi", "xrusd", "gdp", "pop", "crisisJST"]


def main():
    wb = openpyxl.load_workbook(os.path.join(HERE, "raw", "JSTdatasetR6.xlsx"), read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ix = {c: hdr.index(c) for c in COLS}
    rows = []
    for r in it:
        rows.append([r[ix[c]] for c in COLS])
    outp = os.path.join(HERE, "processed", "jst_annual.csv")
    with open(outp, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        for r in rows:
            w.writerow(["" if v is None else v for v in r])
    # coverage report
    cov = {}
    for r in rows:
        iso = r[1]
        ok = all(r[i] not in (None, "") for i in (2, 3, 4, 5))
        if ok:
            cov.setdefault(iso, []).append(r[0])
    for iso, ys in sorted(cov.items()):
        gaps = [y for y in range(ys[0], ys[-1] + 1) if y not in set(ys)]
        print(f"{iso}: {ys[0]}-{ys[-1]} n={len(ys)} gaps={len(gaps)} {gaps[:6]}")
    print("wrote", outp, len(rows))


if __name__ == "__main__":
    main()
